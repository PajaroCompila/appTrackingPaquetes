from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BASE = Path(__file__).resolve().parent
CAPTURAS = BASE / "reporte-mejoras-capturas"
SALIDA = BASE / "Informe de mejoras - Pedidos Bodega.docx"
LOGO = BASE.parent / "frontend" / "public" / "imagenes" / "logo-pajaro-azul.png"

# compact_reference_guide + acabado visual de Pedidos Bodega.
AZUL = "123568"
AZUL_MEDIO = "235A96"
AZUL_CLARO = "E8EEF5"
NARANJA = "F59E0B"
GRIS = "54667A"
GRIS_CLARO = "F4F6F9"
BLANCO = "FFFFFF"
NEGRO = "17212B"
VERDE = "E5F5EC"
ANCHO_UTIL_DXA = 9360


def color(valor: str) -> RGBColor:
    return RGBColor.from_string(valor)


def aplicar_fuente(run, tamano=11, negrita=False, color_texto=NEGRO, cursiva=False):
    run.font.name = "Segoe UI"
    rpr = run._element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:ascii"), "Segoe UI")
    rpr.rFonts.set(qn("w:hAnsi"), "Segoe UI")
    run.font.size = Pt(tamano)
    run.font.bold = negrita
    run.font.italic = cursiva
    run.font.color.rgb = color(color_texto)


def fondo_celda(celda, relleno):
    tcpr = celda._tc.get_or_add_tcPr()
    sombreado = tcpr.find(qn("w:shd"))
    if sombreado is None:
        sombreado = OxmlElement("w:shd")
        tcpr.append(sombreado)
    sombreado.set(qn("w:fill"), relleno)


def margen_celda(celda, arriba=80, abajo=80, inicio=120, fin=120):
    tcpr = celda._tc.get_or_add_tcPr()
    margenes = tcpr.find(qn("w:tcMar"))
    if margenes is None:
        margenes = OxmlElement("w:tcMar")
        tcpr.append(margenes)
    for nombre, valor in (("top", arriba), ("bottom", abajo), ("start", inicio), ("end", fin)):
        nodo = margenes.find(qn(f"w:{nombre}"))
        if nodo is None:
            nodo = OxmlElement(f"w:{nombre}")
            margenes.append(nodo)
        nodo.set(qn("w:w"), str(valor))
        nodo.set(qn("w:type"), "dxa")


def bordes(tabla, color_borde="D6E0EA", tamano="6"):
    tblpr = tabla._tbl.tblPr
    nodo_bordes = tblpr.find(qn("w:tblBorders"))
    if nodo_bordes is None:
        nodo_bordes = OxmlElement("w:tblBorders")
        tblpr.append(nodo_bordes)
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        borde = nodo_bordes.find(qn(f"w:{lado}"))
        if borde is None:
            borde = OxmlElement(f"w:{lado}")
            nodo_bordes.append(borde)
        borde.set(qn("w:val"), "single")
        borde.set(qn("w:sz"), tamano)
        borde.set(qn("w:color"), color_borde)


def geometria_tabla(tabla, anchos):
    tabla.autofit = False
    total = sum(anchos)
    tblpr = tabla._tbl.tblPr
    tblw = tblpr.find(qn("w:tblW"))
    if tblw is None:
        tblw = OxmlElement("w:tblW")
        tblpr.append(tblw)
    tblw.set(qn("w:w"), str(total))
    tblw.set(qn("w:type"), "dxa")
    tblind = tblpr.find(qn("w:tblInd"))
    if tblind is None:
        tblind = OxmlElement("w:tblInd")
        tblpr.append(tblind)
    tblind.set(qn("w:w"), "120")
    tblind.set(qn("w:type"), "dxa")
    grid = tabla._tbl.tblGrid
    for hijo in list(grid):
        grid.remove(hijo)
    for valor in anchos:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(valor))
        grid.append(col)
    for fila in tabla.rows:
        for indice, celda in enumerate(fila.cells):
            valor = anchos[min(indice, len(anchos) - 1)]
            tcw = celda._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tcw is None:
                tcw = OxmlElement("w:tcW")
                celda._tc.get_or_add_tcPr().append(tcw)
            tcw.set(qn("w:w"), str(valor))
            tcw.set(qn("w:type"), "dxa")


def configurar_estilos(doc):
    seccion = doc.sections[0]
    seccion.page_width = Inches(8.5)
    seccion.page_height = Inches(11)
    seccion.top_margin = Inches(1)
    seccion.bottom_margin = Inches(1)
    seccion.left_margin = Inches(1)
    seccion.right_margin = Inches(1)
    seccion.header_distance = Inches(0.492)
    seccion.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Segoe UI"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Segoe UI")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Segoe UI")
    normal.font.size = Pt(11)
    normal.font.color.rgb = color(NEGRO)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    tokens = {
        "Title": (28, AZUL, 0, 8),
        "Subtitle": (14, GRIS, 0, 12),
        "Heading 1": (16, AZUL_MEDIO, 18, 10),
        "Heading 2": (13, AZUL_MEDIO, 14, 7),
        "Heading 3": (12, AZUL, 10, 5),
    }
    for nombre, (tamano, tono, antes, despues) in tokens.items():
        estilo = doc.styles[nombre]
        estilo.font.name = "Segoe UI"
        estilo._element.rPr.rFonts.set(qn("w:ascii"), "Segoe UI")
        estilo._element.rPr.rFonts.set(qn("w:hAnsi"), "Segoe UI")
        estilo.font.size = Pt(tamano)
        estilo.font.bold = nombre != "Subtitle"
        estilo.font.color.rgb = color(tono)
        estilo.paragraph_format.space_before = Pt(antes)
        estilo.paragraph_format.space_after = Pt(despues)
        estilo.paragraph_format.keep_with_next = True

    encabezado = seccion.header.paragraphs[0]
    encabezado.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    aplicar_fuente(encabezado.add_run("PEDIDOS BODEGA  |  MEJORAS REALIZADAS"), 8, True, GRIS)

    pie = seccion.footer.paragraphs[0]
    pie.alignment = WD_ALIGN_PARAGRAPH.CENTER
    aplicar_fuente(pie.add_run("Todos los derechos reservados. 2026© Almacén Pájaro Azul.  ·  Página "), 8, False, GRIS)
    campo = OxmlElement("w:fldSimple")
    campo.set(qn("w:instr"), "PAGE")
    pie._p.append(campo)


def parrafo(doc, texto, tamano=11, negrita=False, tono=NEGRO, centrado=False, despues=6, cursiva=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if centrado else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(despues)
    aplicar_fuente(p.add_run(texto), tamano, negrita, tono, cursiva)
    return p


def agregar_resumen(doc):
    doc.add_heading("Cambios incluidos", level=1)
    parrafo(doc, "Esta lista sirve como guía rápida antes de revisar las capturas.", 10.5, False, GRIS, despues=10)
    filas = [
        ("Acceso", "El fondo llena el lado izquierdo y mantiene visibles las marcas."),
        ("Uso diario", "Botones, enlaces, tarjetas, pestañas y filas responden al mouse y al teclado con movimientos cortos."),
        ("Filtros", "Los almacenes elegidos quedan visibles y se conservan al pasar de Pedidos a Despachados e Historial."),
        ("Pedidos", "Cada artículo aparece en su propia fila. El despacho continúa trabajando con el pedido completo."),
        ("Historial", "La información puede revisarse por pedido o por artículo. Código y descripción permiten abrir el inventario con doble clic."),
        ("Dashboard", "Los pedidos pendientes y entregados de las seis sucursales aparecen en una sola vista."),
        ("Usuarios", "Crear, buscar, editar, activar, desactivar y restablecer contraseña está reunido en una sola pantalla."),
    ]
    tabla = doc.add_table(rows=1, cols=2)
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
    geometria_tabla(tabla, [2160, 7200])
    bordes(tabla)
    for i, texto in enumerate(("Área", "Qué cambió")):
        celda = tabla.rows[0].cells[i]
        fondo_celda(celda, AZUL_CLARO)
        margen_celda(celda)
        aplicar_fuente(celda.paragraphs[0].add_run(texto), 10, True, AZUL)
    for indice, (area, detalle) in enumerate(filas):
        celdas = tabla.add_row().cells
        for celda in celdas:
            celda.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margen_celda(celda, 105, 105, 120, 120)
            fondo_celda(celda, BLANCO if indice % 2 == 0 else GRIS_CLARO)
        aplicar_fuente(celdas[0].paragraphs[0].add_run(area), 10, True, AZUL)
        aplicar_fuente(celdas[1].paragraphs[0].add_run(detalle), 10)
    parrafo(doc, "Cómo leer las capturas", 12, True, AZUL, despues=4)
    parrafo(doc, "Busque el número naranja en la imagen y lea el número igual que aparece debajo.", 10.5, False, GRIS)


def agregar_llamadas(doc, items):
    tabla = doc.add_table(rows=0, cols=2)
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
    for numero, texto in items:
        celdas = tabla.add_row().cells
        celdas[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        celdas[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        margen_celda(celdas[0], 70, 70, 120, 120)
        margen_celda(celdas[1], 70, 70, 120, 120)
        fondo_celda(celdas[0], NARANJA)
        fondo_celda(celdas[1], GRIS_CLARO)
        p_numero = celdas[0].paragraphs[0]
        p_numero.alignment = WD_ALIGN_PARAGRAPH.CENTER
        aplicar_fuente(p_numero.add_run(str(numero)), 11, True, BLANCO)
        aplicar_fuente(celdas[1].paragraphs[0].add_run(texto), 9.5, False, NEGRO)
    geometria_tabla(tabla, [640, 8720])
    bordes(tabla, "D6E0EA", "4")


def pagina_captura(doc, titulo, texto, archivo, items, nota=None):
    doc.add_page_break()
    doc.add_heading(titulo, level=1)
    parrafo(doc, texto, 10.5, False, GRIS, despues=8)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    dibujo = p.add_run().add_picture(str(CAPTURAS / archivo), width=Inches(6.42))
    dibujo._inline.docPr.set("descr", f"Captura real de {titulo.lower()} en Pedidos Bodega")
    parrafo(doc, "Captura tomada en la aplicación en funcionamiento.", 8, False, GRIS, centrado=True, despues=5, cursiva=True)
    agregar_llamadas(doc, items)
    if nota:
        p_nota = doc.add_paragraph()
        p_nota.paragraph_format.space_before = Pt(5)
        p_nota.paragraph_format.space_after = Pt(0)
        aplicar_fuente(p_nota.add_run("Nota: "), 9, True, AZUL)
        aplicar_fuente(p_nota.add_run(nota), 9, False, GRIS)


doc = Document()
configurar_estilos(doc)
doc.core_properties.title = "Mejoras realizadas - Pedidos Bodega"
doc.core_properties.subject = "Capturas reales de las mejoras visibles en la aplicación"
doc.core_properties.author = "Almacén Pájaro Azul"

# Portada: editorial_cover con color y tipografía de la aplicación.
parrafo(doc, "PEDIDOS BODEGA", 10, True, AZUL_MEDIO, centrado=True, despues=36)
if LOGO.exists():
    p_logo = doc.add_paragraph()
    p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_logo.paragraph_format.space_after = Pt(36)
    logo = p_logo.add_run().add_picture(str(LOGO), width=Inches(3.2))
    logo._inline.docPr.set("descr", "Logotipo de Almacén Pájaro Azul")
parrafo(doc, "Mejoras realizadas", 28, True, AZUL, centrado=True, despues=8)
parrafo(doc, "Capturas de la aplicación en funcionamiento", 14, False, GRIS, centrado=True, despues=34)
parrafo(doc, "Cada número naranja señala el lugar exacto del cambio explicado debajo de la imagen.", 11, False, NEGRO, centrado=True, despues=36)
parrafo(doc, "3 de septiembre de 2026", 10, True, GRIS, centrado=True, despues=4)

doc.add_page_break()
agregar_resumen(doc)

pagina_captura(
    doc,
    "Inicio de sesión",
    "El acceso conserva la distribución de dos paneles y aprovecha toda la altura de la pantalla.",
    "01-login.png",
    [
        (1, "El fondo cubre todo el lado izquierdo. La imagen no se estira y las marcas del extremo izquierdo quedan completas."),
        (2, "El campo Usuario mantiene el texto visible y el foco se reconoce con claridad."),
        (3, "El campo Contraseña permite mostrar u ocultar lo escrito sin cambiar el tamaño del formulario."),
        (4, "Iniciar sesión se habilita cuando ambos datos están completos."),
    ],
)

pagina_captura(
    doc,
    "Respuesta de botones y tarjetas",
    "Esta pantalla se usó para revisar los movimientos antes de llevarlos a las pantallas de trabajo.",
    "02-microinteracciones.png",
    [
        (1, "Los botones responden al pasar el mouse, al presionar y al navegar con la tecla Tab."),
        (2, "Las tarjetas que abren otra vista se levantan ligeramente. Las tarjetas informativas permanecen quietas."),
        (3, "Pestañas y filas remarcan la opción sobre la que se trabaja sin mover el resto de la página."),
    ],
    "La pantalla de prueba no consulta ni modifica pedidos, inventario, usuarios, SAP o RetailOne.",
)

pagina_captura(
    doc,
    "Dashboard por sucursal",
    "El dashboard deja los totales y el resultado de cada sucursal en la misma pantalla.",
    "08-dashboard.png",
    [
        (1, "La fecha y la tienda se filtran desde una sola línea de controles."),
        (2, "Los totales de pendientes y entregados mantienen el mismo tamaño y alineación."),
        (3, "Las seis sucursales aparecen en tarjetas parejas. Si una no responde, se indica como No disponible."),
    ],
    "El dashboard está disponible únicamente para usuarios administradores.",
)

pagina_captura(
    doc,
    "Pedidos de bodega",
    "La selección de almacenes queda visible y la tabla mantiene una fila por cada artículo.",
    "03-pedidos-filtros.png",
    [
        (1, "El selector indica cuántos almacenes están activos."),
        (2, "Cada código seleccionado aparece como una etiqueta. La x quita solo ese almacén."),
        (3, "Buscar y Limpiar filtros conservan la misma posición que en las demás pantallas."),
        (4, "Cada artículo ocupa su propia fila con pedido, código, descripción, cantidad, bodega, fecha y vendedor."),
    ],
    "El refresco en segundo plano conserva los filtros y las selecciones de artículos.",
)

pagina_captura(
    doc,
    "Pedidos despachados",
    "Los almacenes elegidos en Pedidos siguen activos al entrar a Despachados.",
    "04-despachados-filtros.png",
    [
        (1, "El resumen mantiene la cantidad de almacenes seleccionados."),
        (2, "BSPS03 y TSPS01 continúan visibles; no fue necesario volver a marcarlos."),
        (3, "Los botones usan el mismo orden, separación y tamaño que en Pedidos e Historial."),
    ],
    "La captura muestra que no había pedidos despachados con esos almacenes y esa fecha.",
)

pagina_captura(
    doc,
    "Historial por pedido",
    "La vista Pedido resume cada entrega en una sola fila y conserva los filtros seleccionados.",
    "05-historial-pedidos.png",
    [
        (1, "El selector conserva los dos almacenes que se eligieron en Pedidos."),
        (2, "Las etiquetas permiten revisar y retirar cada almacén por separado."),
        (3, "Pedido y Artículos cambian la forma de ver el mismo historial."),
        (4, "La tabla por pedido muestra estado, cantidad de artículos, bodegas, fecha, vendedor, usuario y acceso al detalle."),
    ],
)

pagina_captura(
    doc,
    "Historial por artículo",
    "La vista Artículos presenta cada producto en una fila sin perder la relación con su pedido.",
    "06-historial-articulos.png",
    [
        (1, "La pestaña Artículos queda marcada para que se reconozca la vista activa."),
        (2, "La tabla mantiene una línea por artículo y muestra la columna Estado."),
        (3, "Un doble clic sobre el código abre la consulta de inventario."),
        (4, "Un doble clic sobre la descripción abre la misma consulta de inventario."),
    ],
)

pagina_captura(
    doc,
    "Administración de usuarios",
    "La pantalla reúne el alta, la búsqueda y el mantenimiento de cuentas locales.",
    "09-usuarios.png",
    [
        (1, "Crear usuario abre el formulario de nombre, usuario, rol y contraseña. No solicita correo ni obliga a cambiar la contraseña al ingresar."),
        (2, "La búsqueda se puede reducir por nombre o usuario, rol y estado."),
        (3, "La tabla muestra el último acceso y permite reconocer cuentas activas e inactivas."),
        (4, "Editar, activar o desactivar y restablecer contraseña están disponibles en la misma fila."),
    ],
    "Esta pantalla está disponible únicamente para usuarios administradores.",
)

doc.save(SALIDA)
print(SALIDA)
